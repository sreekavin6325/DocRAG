"""
rag_core.py — the retrieval engine for the document-search RAG app.

Fully offline and dependency-light. Indexes your own documents
(PDF / DOCX / TXT / MD), splits them into overlapping chunks, builds a TF-IDF
inverted index (implemented in pure Python — no scikit-learn / SciPy), and
retrieves the most relevant passages for a query by cosine similarity.

Answer generation is pluggable:
  - default  -> "extractive": returns the actual passages from your documents.
  - optional -> an LLM synthesises an answer from those passages, if you supply
                an API key (OpenAI or Anthropic). Nothing is sent anywhere unless
                you explicitly configure a key.
"""

from __future__ import annotations

import os
import re
import math
import pickle
from collections import defaultdict, Counter
from dataclasses import dataclass
from typing import Iterable

SUPPORTED_EXTS = {
    # plain text / markup
    ".txt", ".md", ".markdown", ".log", ".rst",
    ".csv", ".tsv", ".json", ".xml", ".html", ".htm", ".rtf",
    ".py", ".js", ".ts", ".java", ".c", ".cpp", ".cs", ".go", ".rb",
    ".ini", ".cfg", ".yaml", ".yml",
    # office / documents
    ".pdf", ".docx", ".xlsx", ".xlsm", ".xls", ".pptx",
}

# Extensions read as raw text with the generic text reader.
_TEXT_EXTS = {
    ".txt", ".md", ".markdown", ".log", ".rst", ".xml",
    ".py", ".js", ".ts", ".java", ".c", ".cpp", ".cs", ".go", ".rb",
    ".ini", ".cfg", ".yaml", ".yml",
}

TOKEN_RE = re.compile(r"[a-z0-9_']+")

# A compact English stop-word list (enough to cut the noise without a dependency).
STOP_WORDS = {
    "a", "an", "and", "are", "as", "at", "be", "but", "by", "for", "if", "in",
    "into", "is", "it", "its", "no", "not", "of", "on", "or", "such", "that",
    "the", "their", "then", "there", "these", "they", "this", "to", "was",
    "will", "with", "we", "you", "your", "i", "me", "my", "he", "she", "him",
    "her", "from", "have", "has", "had", "do", "does", "did", "can", "could",
    "would", "should", "about", "which", "who", "whom", "what", "when", "where",
    "how", "all", "any", "been", "being", "were", "am", "so", "than", "too",
    "very", "just", "also", "more", "most", "some", "up", "out", "over", "our",
}


# --------------------------------------------------------------------------- #
# Document loading
# --------------------------------------------------------------------------- #
def read_text_file(path: str) -> str:
    for enc in ("utf-8", "utf-8-sig", "latin-1"):
        try:
            with open(path, "r", encoding=enc) as fh:
                return fh.read()
        except (UnicodeDecodeError, UnicodeError):
            continue
    with open(path, "rb") as fh:
        return fh.read().decode("utf-8", errors="replace")


def read_pdf(path: str) -> str:
    from pypdf import PdfReader

    reader = PdfReader(path)
    parts = []
    for page in reader.pages:
        try:
            parts.append(page.extract_text() or "")
        except Exception:
            parts.append("")
    return "\n".join(parts)


def read_docx(path: str) -> str:
    import docx

    doc = docx.Document(path)
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append("\t".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def read_xlsx(path: str) -> str:
    """Read every sheet of an .xlsx/.xlsm workbook as 'sheet | cell1 | cell2 ...' rows."""
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    parts = []
    for ws in wb.worksheets:
        parts.append(f"# SHEET: {ws.title}")
        for row in ws.iter_rows(values_only=True):
            cells = [str(c) for c in row if c is not None and str(c).strip()]
            if cells:
                parts.append(" | ".join(cells))
    wb.close()
    return "\n".join(parts)


def read_xls(path: str) -> str:
    """Read a legacy .xls workbook via xlrd."""
    import xlrd

    book = xlrd.open_workbook(path)
    parts = []
    for sheet in book.sheets():
        parts.append(f"# SHEET: {sheet.name}")
        for r in range(sheet.nrows):
            cells = [str(c) for c in sheet.row_values(r) if str(c).strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def read_csv(path: str) -> str:
    """Read CSV/TSV, joining each row's cells with ' | '."""
    import csv

    delim = "\t" if path.lower().endswith(".tsv") else ","
    rows = []
    for enc in ("utf-8-sig", "utf-8", "latin-1"):
        try:
            with open(path, "r", encoding=enc, newline="") as fh:
                for row in csv.reader(fh, delimiter=delim):
                    cells = [c for c in row if c and c.strip()]
                    if cells:
                        rows.append(" | ".join(cells))
            return "\n".join(rows)
        except (UnicodeDecodeError, UnicodeError):
            rows = []
            continue
    return ""


def read_pptx(path: str) -> str:
    from pptx import Presentation

    prs = Presentation(path)
    parts = []
    for i, slide in enumerate(prs.slides, 1):
        parts.append(f"# SLIDE {i}")
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    line = "".join(run.text for run in para.runs)
                    if line.strip():
                        parts.append(line)
            if shape.has_table:
                for row in shape.table.rows:
                    cells = [c.text for c in row.cells if c.text.strip()]
                    if cells:
                        parts.append(" | ".join(cells))
    return "\n".join(parts)


def read_html(path: str) -> str:
    """Strip tags/scripts/styles and return visible text."""
    raw = read_text_file(path)
    raw = re.sub(r"(?is)<(script|style).*?</\1>", " ", raw)
    text = re.sub(r"(?s)<[^>]+>", " ", raw)
    text = re.sub(r"&nbsp;", " ", text)
    text = re.sub(r"&amp;", "&", text)
    text = re.sub(r"&lt;", "<", text).replace("&gt;", ">")
    return re.sub(r"[ \t]*\n[ \t]*", "\n", re.sub(r"[ \t]+", " ", text)).strip()


def read_json(path: str) -> str:
    """Flatten JSON into readable 'key: value' lines."""
    import json as _json

    try:
        data = _json.loads(read_text_file(path))
    except Exception:
        return read_text_file(path)

    lines: list[str] = []

    def walk(obj, prefix=""):
        if isinstance(obj, dict):
            for k, v in obj.items():
                walk(v, f"{prefix}{k}.")
        elif isinstance(obj, list):
            for i, v in enumerate(obj):
                walk(v, f"{prefix}{i}.")
        else:
            lines.append(f"{prefix.rstrip('.')}: {obj}")

    walk(data)
    return "\n".join(lines)


def read_rtf(path: str) -> str:
    """Very small RTF-to-text stripper (control words + groups removed)."""
    raw = read_text_file(path)
    raw = re.sub(r"\\'[0-9a-fA-F]{2}", " ", raw)
    raw = re.sub(r"\\[a-zA-Z]+-?\d* ?", " ", raw)
    raw = raw.replace("{", " ").replace("}", " ")
    return re.sub(r"[ \t]+", " ", raw).strip()


def load_document(path: str) -> str:
    ext = os.path.splitext(path)[1].lower()
    if ext in _TEXT_EXTS:
        return read_text_file(path)
    if ext == ".pdf":
        return read_pdf(path)
    if ext == ".docx":
        return read_docx(path)
    if ext in (".xlsx", ".xlsm"):
        return read_xlsx(path)
    if ext == ".xls":
        return read_xls(path)
    if ext in (".csv", ".tsv"):
        return read_csv(path)
    if ext == ".pptx":
        return read_pptx(path)
    if ext in (".html", ".htm"):
        return read_html(path)
    if ext == ".json":
        return read_json(path)
    if ext == ".rtf":
        return read_rtf(path)
    raise ValueError(f"Unsupported file type: {ext}")


def discover_files(root: str) -> list[str]:
    """Return every supported file under `root` (file or directory)."""
    if os.path.isfile(root):
        return [root] if os.path.splitext(root)[1].lower() in SUPPORTED_EXTS else []
    found = []
    for dirpath, _dirs, files in os.walk(root):
        for name in files:
            if os.path.splitext(name)[1].lower() in SUPPORTED_EXTS:
                found.append(os.path.join(dirpath, name))
    return sorted(found)


# --------------------------------------------------------------------------- #
# Tokenisation (unigrams + bigrams)
# --------------------------------------------------------------------------- #
def tokenize(text: str) -> list[str]:
    words = [w for w in TOKEN_RE.findall(text.lower()) if len(w) > 1 and w not in STOP_WORDS]
    terms = list(words)
    # bigrams improve phrase matching (e.g. "recovery codes")
    terms.extend(f"{words[i]} {words[i + 1]}" for i in range(len(words) - 1))
    return terms


# --------------------------------------------------------------------------- #
# Chunking
# --------------------------------------------------------------------------- #
def chunk_text(text: str, chunk_size: int = 900, overlap: int = 150) -> list[str]:
    """Split text into overlapping chunks of up to `chunk_size` words.

    Line-aware: line breaks are preserved inside each chunk, so spreadsheet /
    CSV rows and list items stay on their own line (the answer composer can then
    isolate a single matching row instead of returning a whole sheet). A single
    very long line is split on word boundaries.
    """
    text = re.sub(r"\r\n?", "\n", text)
    lines: list[str] = []
    for raw_line in text.split("\n"):
        line = raw_line.strip()
        if not line:
            continue
        words = line.split()
        if len(words) <= chunk_size:
            lines.append(line)
        else:  # break an over-long line into word-sized pieces
            for i in range(0, len(words), chunk_size):
                lines.append(" ".join(words[i : i + chunk_size]))

    chunks: list[str] = []
    buf: list[str] = []
    buf_words = 0

    for line in lines:
        w = len(line.split())
        if buf_words + w > chunk_size and buf:
            chunks.append("\n".join(buf))
            # carry a trailing overlap (whole lines) into the next chunk
            tail: list[str] = []
            tw = 0
            for b in reversed(buf):
                tail.insert(0, b)
                tw += len(b.split())
                if tw >= overlap:
                    break
            buf, buf_words = list(tail), sum(len(b.split()) for b in tail)
        buf.append(line)
        buf_words += w

    if buf:
        chunks.append("\n".join(buf))
    return [c for c in chunks if c.strip()]


# --------------------------------------------------------------------------- #
# Index
# --------------------------------------------------------------------------- #
@dataclass
class Chunk:
    doc_id: int
    source: str
    chunk_index: int
    text: str


@dataclass
class SearchHit:
    score: float
    source: str
    chunk_index: int
    text: str


class RagIndex:
    """Pure-Python TF-IDF inverted index with cosine-similarity retrieval."""

    def __init__(self, chunk_size: int = 900, overlap: int = 150):
        self.chunk_size = chunk_size
        self.overlap = overlap
        self.chunks: list[Chunk] = []
        self.sources: list[str] = []
        # learned at build time
        self.idf: dict[str, float] = {}
        self.postings: dict[str, list[tuple[int, float]]] = {}  # term -> [(chunk_id, weight)]
        self.norms: list[float] = []

    # ---- building -------------------------------------------------------- #
    def build(self, paths: Iterable[str], progress=None) -> dict:
        files: list[str] = []
        for p in paths:
            files.extend(discover_files(p))
        files = sorted(set(files))

        self.chunks = []
        self.sources = []
        errors: list[str] = []
        total = len(files)

        for doc_id, path in enumerate(files):
            if progress:
                progress(doc_id, total, f"Reading {os.path.basename(path)}")
            try:
                raw = load_document(path)
            except Exception as exc:  # noqa: BLE001
                errors.append(f"{path}: {exc}")
                continue
            self.sources.append(path)
            for ci, chunk in enumerate(chunk_text(raw, self.chunk_size, self.overlap)):
                self.chunks.append(Chunk(doc_id, path, ci, chunk))

        self.idf, self.postings, self.norms = {}, {}, []
        if not self.chunks:
            return {"files": len(self.sources), "chunks": 0, "errors": errors}

        if progress:
            progress(total, total, "Building index")

        # first pass: term frequencies per chunk + document frequencies
        tf_per_chunk: list[Counter] = []
        df: Counter = Counter()
        for chunk in self.chunks:
            counts = Counter(tokenize(chunk.text))
            tf_per_chunk.append(counts)
            for term in counts:
                df[term] += 1

        n = len(self.chunks)
        self.idf = {t: math.log((1 + n) / (1 + d)) + 1.0 for t, d in df.items()}

        # second pass: build normalised tf-idf vectors as an inverted index
        postings: dict[str, list[tuple[int, float]]] = defaultdict(list)
        norms = [0.0] * n
        for cid, counts in enumerate(tf_per_chunk):
            sq = 0.0
            weights = {}
            for term, tf in counts.items():
                w = (1.0 + math.log(tf)) * self.idf[term]  # sublinear tf
                weights[term] = w
                sq += w * w
            norm = math.sqrt(sq) or 1.0
            norms[cid] = norm
            for term, w in weights.items():
                postings[term].append((cid, w / norm))  # store L2-normalised weight
        self.postings = dict(postings)
        self.norms = norms

        return {"files": len(self.sources), "chunks": n, "errors": errors}

    # ---- querying -------------------------------------------------------- #
    @property
    def ready(self) -> bool:
        return bool(self.postings)

    def search(self, query: str, top_k: int = 5) -> list[SearchHit]:
        if not self.ready or not query.strip():
            return []

        q_counts = Counter(tokenize(query))
        if not q_counts:
            return []

        q_weights = {}
        sq = 0.0
        for term, tf in q_counts.items():
            if term not in self.idf:
                continue
            w = (1.0 + math.log(tf)) * self.idf[term]
            q_weights[term] = w
            sq += w * w
        if not q_weights:
            return []
        q_norm = math.sqrt(sq) or 1.0

        scores: dict[int, float] = defaultdict(float)
        for term, qw in q_weights.items():
            qn = qw / q_norm
            for cid, cw in self.postings.get(term, ()):  # cw already normalised
                scores[cid] += qn * cw

        ranked = sorted(scores.items(), key=lambda kv: kv[1], reverse=True)[:top_k]
        hits = []
        for cid, score in ranked:
            if score <= 0:
                continue
            c = self.chunks[cid]
            hits.append(SearchHit(float(score), c.source, c.chunk_index, c.text))
        return hits

    # ---- persistence ----------------------------------------------------- #
    def save(self, path: str) -> None:
        with open(path, "wb") as fh:
            pickle.dump(
                {
                    "chunk_size": self.chunk_size,
                    "overlap": self.overlap,
                    "chunks": [(c.doc_id, c.source, c.chunk_index, c.text) for c in self.chunks],
                    "sources": self.sources,
                    "idf": self.idf,
                    "postings": self.postings,
                    "norms": self.norms,
                },
                fh,
            )

    @classmethod
    def load(cls, path: str) -> "RagIndex":
        with open(path, "rb") as fh:
            data = pickle.load(fh)
        idx = cls(data.get("chunk_size", 900), data.get("overlap", 150))
        idx.chunks = [Chunk(*t) for t in data["chunks"]]
        idx.sources = data["sources"]
        idx.idf = data["idf"]
        idx.postings = data["postings"]
        idx.norms = data["norms"]
        return idx


# --------------------------------------------------------------------------- #
# Answer composition
# --------------------------------------------------------------------------- #
def _segment_split(text: str) -> list[str]:
    """Split a passage into candidate answer segments.

    Splits on both sentence boundaries AND line breaks, so prose yields
    sentences while spreadsheet / CSV rows yield individual rows. This lets the
    composer return the single matching table row instead of the whole sheet.
    """
    segments: list[str] = []
    for line in text.split("\n"):
        line = line.strip()
        if not line or line.startswith("#"):  # skip sheet/slide markers
            continue
        for sent in re.split(r"(?<=[.!?])\s+", line):
            sent = sent.strip()
            if sent:
                segments.append(sent)
    return segments


def extractive_answer(query: str, hits: list[SearchHit], max_sentences: int = 4) -> str:
    """Pick the segments from the retrieved passages most relevant to the query.

    Fully offline; no model involved. This is the default 'generation' step.
    """
    if not hits:
        return "No relevant passages found in the indexed documents."

    terms = {w for w in re.findall(r"\w+", query.lower()) if len(w) > 2}
    scored: list[tuple[float, str]] = []
    for hit in hits:
        for sent in _segment_split(hit.text):
            words = re.findall(r"\w+", sent.lower())
            if not words:
                continue
            overlap = sum(1 for w in words if w in terms)
            if overlap == 0:
                continue
            # reward denser matches; mild length penalty keeps whole-sheet dumps out
            score = overlap * 2 + hit.score - 0.002 * len(words)
            scored.append((score, sent))

    scored.sort(key=lambda x: x[0], reverse=True)
    seen = set()
    picked = []
    for _score, sent in scored:
        key = sent.lower()
        if key in seen:
            continue
        seen.add(key)
        picked.append(sent)
        if len(picked) >= max_sentences:
            break

    if not picked:
        return hits[0].text[:800]
    return " ".join(picked)


def build_prompt(query: str, hits: list[SearchHit], max_context_chars: int = 6000) -> str:
    """Compose a RAG prompt from retrieved context for an optional LLM step.

    The context is capped at `max_context_chars` (roughly ~1.5k tokens) so a
    single request stays within free-tier per-minute token limits. Passages are
    added whole until the budget runs out, then the last one is truncated.
    """
    blocks = []
    used = 0
    for i, hit in enumerate(hits, 1):
        header = f"[{i}] (source: {os.path.basename(hit.source)})\n"
        remaining = max_context_chars - used
        if remaining <= len(header) + 50:
            break
        text = hit.text
        if len(header) + len(text) > remaining:
            text = text[: remaining - len(header)].rsplit(" ", 1)[0] + " …"
        block = header + text
        blocks.append(block)
        used += len(block) + 2
    context = "\n\n".join(blocks)
    return (
        "Answer the question using ONLY the context below. "
        "Cite sources with their [number]. If the answer is not in the context, "
        "say you don't know.\n\n"
        f"Context:\n{context}\n\n"
        f"Question: {query}\n\nAnswer:"
    )


def llm_answer(query: str, hits: list[SearchHit], config: dict) -> str:
    """Optional LLM synthesis. `config` = {provider, api_key, model}."""
    provider = config.get("provider", "").lower()
    api_key = config.get("api_key", "")
    if not provider or not api_key:
        raise ValueError("LLM not configured")

    prompt = build_prompt(query, hits)

    if provider == "openai":
        from openai import OpenAI

        client = OpenAI(api_key=api_key)
        model = config.get("model") or "gpt-4o-mini"
        resp = client.chat.completions.create(
            model=model,
            messages=[{"role": "user", "content": prompt}],
            temperature=0.1,
        )
        return resp.choices[0].message.content.strip()

    if provider == "anthropic":
        import anthropic

        client = anthropic.Anthropic(api_key=api_key)
        model = config.get("model") or "claude-3-5-haiku-latest"
        resp = client.messages.create(
            model=model,
            max_tokens=1024,
            messages=[{"role": "user", "content": prompt}],
        )
        return "".join(b.text for b in resp.content if b.type == "text").strip()

    if provider == "groq":
        # Groq's OpenAI-compatible endpoint, called via the standard library so
        # no extra package is bundled into the exe.
        model = config.get("model") or "llama-3.3-70b-versatile"
        return _openai_compatible_chat(
            url="https://api.groq.com/openai/v1/chat/completions",
            api_key=api_key,
            model=model,
            prompt=prompt,
        )

    raise ValueError(f"Unknown provider: {provider}")


def _openai_compatible_chat(url: str, api_key: str, model: str, prompt: str) -> str:
    """POST a single-message chat completion to an OpenAI-compatible endpoint
    (Groq, OpenAI, local servers) using only the standard library."""
    import re as _re
    import time as _time
    import json as _json
    import urllib.request
    import urllib.error

    body = _json.dumps(
        {
            "model": model,
            "messages": [{"role": "user", "content": prompt}],
            "temperature": 0.1,
        }
    ).encode("utf-8")
    req = urllib.request.Request(
        url,
        data=body,
        headers={
            "Authorization": f"Bearer {api_key}",
            "Content-Type": "application/json",
            # a browser-style UA avoids Cloudflare's default-urllib block (error 1010)
            "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) DocFinderRAG/1.0",
            "Accept": "application/json",
        },
        method="POST",
    )

    for attempt in range(2):  # one automatic retry on a rate-limit
        try:
            with urllib.request.urlopen(req, timeout=60) as resp:
                data = _json.loads(resp.read().decode("utf-8"))
            return data["choices"][0]["message"]["content"].strip()
        except urllib.error.HTTPError as exc:
            detail = exc.read().decode("utf-8", "replace")
            if exc.code == 429 and attempt == 0:
                m = _re.search(r"try again in ([0-9.]+)s", detail)
                wait = min(float(m.group(1)) + 0.5, 12.0) if m else 6.0
                _time.sleep(wait)
                continue
            raise RuntimeError(f"HTTP {exc.code}: {detail[:200]}") from exc


def answer_query(query: str, index: RagIndex, top_k: int = 5, llm_config: dict | None = None):
    """Return (answer_text, hits). Uses the LLM if configured, else extractive."""
    hits = index.search(query, top_k=top_k)
    if llm_config and llm_config.get("provider") and llm_config.get("api_key"):
        try:
            return llm_answer(query, hits, llm_config), hits
        except Exception as exc:  # noqa: BLE001
            fallback = extractive_answer(query, hits)
            return f"(LLM unavailable: {exc}. Showing extracted passages.)\n\n{fallback}", hits
    return extractive_answer(query, hits), hits
